from pathlib import Path
import json, re, datetime, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).parent
out=root/'outputs'
run=Path(r'\\wsl.localhost\Ubuntu-24.04\home\aliwhoami\parg\runs\20261016_132012')
status=json.loads((run/'status.json').read_text())
stamp=datetime.datetime.now().astimezone().isoformat(timespec='seconds')
paragraph=f'Status captured {stamp}. Run state: {status["state"]}. Current stage: {status.get("current_stage") or "none"}. Completed stages: '+', '.join(x['stage'] for x in status['stages'] if x['returncode']==0)+'.'
if (run/'screen_progress.json').exists():
    sp=json.loads((run/'screen_progress.json').read_text())
    paragraph+=f' At this snapshot, screening had processed {sp["processed"]:,} entries and scored {sp["scored"]:,} successfully.'
text=(out/'PARG_Methodology.md').read_text(encoding='utf-8').replace('{{STATUS}}',paragraph)
(out/'PARG_Methodology.md').write_text(text,encoding='utf-8')
(out/'methodology_status_snapshot.json').write_text(json.dumps({'captured':stamp,'status':status},indent=2),encoding='utf-8')
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.8); sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.08; normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(22); doc.styles['Title'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',14),('Heading 2',12)]:
    doc.styles[name].font.size=Pt(size); doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(12)
    doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles[name].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Page ').font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='Computational methodology for PARG ligand modelling and docking'
doc.core_properties.author='Muhammad Ali Khan'
doc.core_properties.subject='Methods and reproducibility record'

def add_table(lines):
    cells=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    cells=[row for row in cells if not all(re.match(r'^[-: ]+$',v) for v in row)]
    n=len(cells[0]); table=doc.add_table(rows=0,cols=n); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths={6:[.48,1.03,1.02,1.02,1.0,1.02],3:[1.3,3.6,1.67]}.get(n,[6.57/n]*n)
    if n==3 and cells[0][2]=='Cases': widths=[1.35,4.67,.55]
    if n==3 and cells[0][0]=='Software': widths=[1.5,2.2,2.87]
    if n==6: widths=[.5,1.10,1.10,1.1,1.1,1.67]
    for i,w in enumerate(widths): table.columns[i].width=Inches(w)
    for j,row in enumerate(cells):
        tr=table.add_row()
        cant=OxmlElement('w:cantSplit'); tr._tr.get_or_add_trPr().append(cant)
        if j==0:
            repeat=OxmlElement('w:tblHeader'); tr._tr.get_or_add_trPr().append(repeat)
        for i,value in enumerate(row):
            cell=tr.cells[i]; cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4); p.paragraph_format.line_spacing=1
            rr=p.add_run(value); rr.font.size=Pt(9); rr.bold=j==0
            if n==6 or (n==3 and i==2 and cells[0][2]=='Cases'): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            if j==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6'); pr.append(sh)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=text.splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i]); i+=1
        add_table(block); continue
    if line.startswith('# '): doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '): doc.add_paragraph(line[4:],'Heading 2')
    else: doc.add_paragraph(line)
    i+=1
path=out/'PARG_Methodology.docx'; doc.save(path)
shutil.copy2(path,run/'PARG_Methodology.docx')
shutil.copy2(out/'PARG_Methodology.md',run/'PARG_Methodology.md')
shutil.copy2(run/'water_selection_audit.json',out/'water_selection_audit.json')
print(path); print('WORDS',len(text.split())); print(paragraph)
